import clevercsv 
import pandas as pd
import json
import re
from docx import Document
from typing import Dict

def print_data(df: pd.DataFrame):
    print("DATA")
    for key in df.keys():
        value = df[key].values[0]
        if key != "Comments":
            print(f"{key}: {value}")
        else: 
            print("Comments:")
            for field in value.split(";"):
                print(f"- {field}")

def create_fields(
    df: pd.DataFrame, mapping: Dict[str, str], regex: Dict[str, any], fields: Dict[str, any]
) -> Dict[str, any]:
    def add_to_fields(key: str, value: str):
        """ Adds key value pair to data-fields"""
        if key in fields:
            # If entry in data-fields is list, parse list
            if isinstance(fields[key], list):
                # Get regex-expression for entries using key
                regex_exp = regex[key]["regex"]
                for x in value.split(", "):
                    try:
                        res = re.search(regex_exp, x)
                        obj = {}
                        for index, name in enumerate(regex[key]["fields"]):
                            obj[name] = res.group(index+1)
                    except:
                        print("parsing regex failed for inp: ", x)
                    fields[key].append(obj)
            # Directly add value
            else:
                fields[key] = value
        # If not directly found, check mapping
        elif key in mapping:
            fields[mapping[key]] = value

    def get_key_val_pair(entry: str):
        """ Gets key-value pairs from list entries (f.e. Comments) """
        key = entry.split(":")[0].lstrip().rstrip()
        val = entry.split(":")[1].lstrip().rstrip()
        return key, val

    for key in df.keys():
        value = df[key].values[0]
        # If not list-type (f.e. Comments) add directly
        if key != "Comments":
            add_to_fields(key, value)
        # For comments, proceed to split list-elements (semicolon-separated)
        else:
            for entry in value.split(";"):
                if ":" in entry:
                    key, val = get_key_val_pair(entry)
                    add_to_fields(key, val)
    return fields

def add_value_paragraph(par, tag, val): 
    """ Adds value to paragraph. Parses value for signal word (f.e. signiture) """
    # Check if value contains a signal-word.
    result = re.search(r"\{(.*)\}", val)
    if result is not None:
        # Get signal-word from regex-match
        rep = result.group(1)
        # For signiture, add signiture-immage and replace tag with ""
        if rep in "signiture":
            p = par.insert_paragraph_before("")
            r = p.add_run()
            r.add_picture("resources/default_sign.png")
            val = ""
    # Replace tag by value
    par.text = par.text.replace(tag, str(val)) 

def edit_paragraph(par, fields):
    """ 
    Edits paragraph by checking for tags and eventually replacing tag with entry
    """
    result = re.search(r"\[(.*)\]", par.text)
    if result is not None:
        key = result.group(1)
        if key in fields:
            tag = "[" + key + "]"
            add_value_paragraph(par, tag, str(fields[key]))
    return par

def edit_paragraphs(doc, fields):
    """ Iterates over all paragraphs in search for tags """
    for par in doc.paragraphs:
        par = edit_paragraph(par, fields)

def edit_tables(doc, fields):
    """ Iterates over all tables in search for tags """
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par = edit_paragraph(par, fields)

def edit_list_tables(doc, fields, defaults):
    """ 
    Iterates over tables in search for list-tags (adding rows per matching entry
    """
    def parse_value(val): 
        """ 
        Checks of value has a replacement (f.e. convert "{yes}" to "✓", but keep
            "20 ml") 
        """
        # Load replacements
        with open("resources/replacements.json") as f:
            replacements = json.load(f)
        # Check for signal-word
        result = re.search(r"\{(.*)\}", val)
        # If found return replacement
        if result is not None:
            rep = result.group(1)
            if rep in replacements:
                return replacements[rep]
        # Otherwise, return original
        return val

    def get_tag(par, remove=False):
        """ 
        Gets (potential) tag from input. If remove is True, remove tag (replace
        with empty string)
        """
        # Search list-tag (<tag>)
        result = re.search(r"\<(.*)\>", par.text)
        # If found return tag
        if result is not None:
            key = result.group(1)
            # If `remove` is True, replace tag with empty string
            if remove:
                tag = "<" + key + ">"
                par.text = par.text.replace(tag, "") 
            return key
        # Otherwise, return None
        return None

    def get_table_infos(row):
        """ Returns first tag (main idenitier) and tags for cells """
        # Get main-tag from first cell
        tag = get_tag(row.cells[0].paragraphs[0])
        # If firt cell has no tag, omit
        if tag is None:
            return None, None
        # Otherwise, load all other tag (for followring cells) too.
        tags = []
        for i in range(0, len(row.cells)):
            tags.append(get_tag(row.cells[i].paragraphs[0], remove=True))
        return tag, tags

    def edit_table(table, tag, tags, value):
        # Add row to table with given information
        row = table.add_row().cells
        name = tags[0]
        for i, x in enumerate(tags):
            # local scope
            if x in value:
                row[i].text = parse_value(value[x])
            # global scope
            elif x in fields:
                row[i].text = parse_value(fields[x])
            # local scope defaults
            elif value[name] in defaults and x in defaults[value[name]]:
                row[i].text = parse_value(defaults[value[name]][x])
            # global scope defaults
            elif x in defaults:
                row[i].text = parse_value(defaults[x])
            else:
                print(f"For tag {value[name]} {x} not found")

    # Iteratre over all tables in search for list-tags
    for table in doc.tables:
        # Get infos from first row
        row = table.rows[0]
        tag, tags = get_table_infos(row)
        if tag is not None and tag in fields:
            # If no information but tag exists, add "None"-row
            if len(fields[tag]) == 0:
                row = table.add_row().cells
                row[0].text = "None"
                for i in range(1, len(tags)):
                    row[i].text = "---"
            # Otherwise add table-row for each entry
            else:
                for entry in fields[tag]:
                    edit_table(table, tag, tags, entry)

# Read input
df = clevercsv.read_dataframe("input/test_en.csv")
# Print original input
print_data(df)

with open("resources/mapping.json") as f:
    mapping = json.load(f)
with open("resources/regex.json") as f:
    regex = json.load(f)
with open("resources/fields.json") as f:
    fields = json.load(f)
with open("resources/defaults.json") as f:
    defaults = json.load(f)

# Parse data-fields from input and print
fields = create_fields(df, mapping, regex, fields)
print("FIELDS:")
for key, value in fields.items():
    print(f"{key}: {value}")

# Load template 
doc = Document("templates/Template_a.docx")
edit_paragraphs(doc, fields)
edit_tables(doc, fields)
edit_list_tables(doc, fields, defaults)

# Save document
doc.save("output/output.docx")
